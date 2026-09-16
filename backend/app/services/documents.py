from __future__ import annotations

import hashlib
import re
from pathlib import Path

from app.models import DocumentClass

JD_FILENAME_HINTS = ("jd", "job_desc", "jobdescription", "job-description")
CV_FILENAME_HINTS = ("cv", "resume")

JD_CONTENT_HINTS = (
    "we are looking for", "responsibilities", "requirements", "about the role",
    "job description", "reporting to",
)
CV_CONTENT_HINTS = (
    "curriculum vitae", "professional experience", "work experience", "education",
    "objective", "references available",
)

COMMAND_VOCABULARY = ("PROCEED", "PAUSE", "REVISE WEIGHTS", "SKIP INTERVIEW")

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff"}


def normalize_to_text(file_path: str) -> str:
    """Best-effort conversion of a fetched file to plain text. Reqruit.ai's
    required DOCX rendition is produced by `render_docx`; this text is only
    for classification, dedup, and the mock extraction heuristics."""
    path = Path(file_path)
    suffix = path.suffix.lower()
    if suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".docx":
        return _read_docx(path)
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix in IMAGE_EXTENSIONS:
        return _ocr_image(path)
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_docx(path: Path) -> str:
    try:
        import docx
    except ImportError:
        return ""
    document = docx.Document(str(path))
    return "\n".join(p.text for p in document.paragraphs)


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _ocr_image(path: Path) -> str:
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        return ""
    try:
        return pytesseract.image_to_string(Image.open(path))
    except Exception:
        return ""


def render_docx(text: str, dest_path: str) -> str:
    """Produces the DOCX rendition Reqruit.ai requires as input (FRD §5.3).
    Falls back to a plain-text file with a .docx name if python-docx isn't
    installed, so the pipeline still runs end-to-end in a minimal mock setup."""
    path = Path(dest_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        import docx
    except ImportError:
        path.write_text(text, encoding="utf-8")
        return str(path)
    document = docx.Document()
    for line in text.splitlines() or [""]:
        document.add_paragraph(line)
    document.save(str(path))
    return str(path)


def classify(text: str, filename: str = "", subject: str = "") -> DocumentClass:
    combined_signal = f"{subject}\n{text}".strip()
    if _is_command(combined_signal):
        return DocumentClass.COMMAND

    lowered_name = filename.lower()
    lowered_text = text.lower()

    if any(h in lowered_name for h in JD_FILENAME_HINTS):
        return DocumentClass.JD
    if any(h in lowered_name for h in CV_FILENAME_HINTS):
        return DocumentClass.CV

    jd_hits = sum(1 for h in JD_CONTENT_HINTS if h in lowered_text)
    cv_hits = sum(1 for h in CV_CONTENT_HINTS if h in lowered_text)
    if jd_hits == 0 and cv_hits == 0:
        return DocumentClass.UNRECOGNIZED
    return DocumentClass.JD if jd_hits >= cv_hits else DocumentClass.CV


def _is_command(text: str) -> bool:
    normalized = re.sub(r"\s+", " ", text).strip().upper()
    return normalized in COMMAND_VOCABULARY


def content_dedup_key(role_id: str, text: str, kind: str) -> str:
    """Hashes normalized text so the same CV/JD pulled twice — even via two
    different connectors — is only screened once (CV-3, JD-6). Whitespace is
    collapsed first since re-conversion of the same file can shift spacing.
    `kind` ("jd"/"cv") keeps a JD and a CV that coincidentally normalize to
    the same text (e.g. two mock fixtures) from colliding on one key."""
    normalized = re.sub(r"\s+", " ", text).strip().lower()
    digest = hashlib.sha256(normalized.encode("utf-8")).hexdigest()
    return f"content:{kind}:{role_id}:{digest}"


def pull_dedup_key(source: str, item_id: str) -> str:
    """Prevents the same provider item from producing a second InboundItem
    when pull() is re-run and nothing new has arrived (CN-5)."""
    return f"pulled:{source}:{item_id}"
